"""Inspecciona el Excel oficial y genera CSV + SQL mediante un mapeo explícito."""
import argparse
import csv
import datetime as dt
import hashlib
import json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
FIELDS = ('nombre_biblioteca', 'tipo_biblioteca', 'departamento', 'provincia', 'distrito', 'entidad_responsable')
REQUIRED = FIELDS[:-1]
LIMITS = dict(nombre_biblioteca=500, tipo_biblioteca=250, departamento=150, provincia=150, distrito=150, entidad_responsable=500)

def clean(value):
    if value is None:
        return None
    if isinstance(value, (dt.date, dt.datetime)):
        return value.isoformat()
    result = str(value).strip()
    return result or None

def literal(value):
    return 'NULL' if value is None else "N'" + str(value).replace("'", "''") + "'"

def insert_lines(table, columns, rows):
    result = []
    for offset in range(0, len(rows), 250):
        chunk = rows[offset:offset + 250]
        result.append(f"INSERT INTO dbo.{table} ({', '.join(columns)}) VALUES\n" + ',\n'.join(
            '(' + ', '.join(str(v) if isinstance(v, int) else literal(v) for v in row) + ')' for row in chunk
        ) + ';')
    return result

def prepare(excel, mapping, output):
    cfg = json.loads(mapping.read_text(encoding='utf-8'))
    wb = openpyxl.load_workbook(excel, read_only=True, data_only=True)
    try:
        if not cfg.get('sheet') or cfg['sheet'] not in wb.sheetnames:
            raise ValueError(f"Configura sheet en mapping.json. Hojas disponibles: {wb.sheetnames}")
        header_row = cfg.get('header_row')
        if not isinstance(header_row, int) or header_row < 1:
            raise ValueError('header_row debe ser un entero positivo.')
        ws = wb[cfg['sheet']]
        header = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
        headers = [clean(x) for x in header]
        present = [x for x in headers if x]
        if len(present) != len(set(present)):
            raise ValueError('La fila de encabezado contiene nombres duplicados: revisa header_row.')
        cols = cfg.get('columns', {})
        index = {}
        for field in FIELDS:
            source = cols.get(field)
            if not source and field in REQUIRED:
                raise ValueError(f'Falta mapear {field}; usa --inspect para ver los encabezados.')
            if source and source not in headers:
                raise ValueError(f'No existe la columna {source!r}. Respeta el encabezado exacto.')
            if source:
                index[field] = headers.index(source)
        source_sha = hashlib.sha256(excel.read_bytes()).hexdigest()
        records = []
        for row_number, values in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            if not any(clean(v) for v in values):
                continue
            rec = {field: clean(values[index[field]]) if field in index else None for field in FIELDS}
            if not rec['nombre_biblioteca']:
                raise ValueError(f'Fila {row_number}: nombre vacío. Revisa encabezado, subtítulos o totales; no se descartan filas silenciosamente.')
            for field in REQUIRED[1:]:
                rec[field] = rec[field] or 'Sin dato'
            for field, limit in LIMITS.items():
                if rec[field] and len(rec[field].encode('utf-16-le')) // 2 > limit:
                    raise ValueError(f'Fila {row_number}: {field} excede el tamaño SQL {limit}.')
            rec['registro_id'] = hashlib.sha256(f'{source_sha}|{ws.title}|{row_number}'.encode()).hexdigest()
            rec['hoja_origen'] = ws.title
            rec['fila_origen'] = row_number
            rec['datos_originales'] = json.dumps({h: clean(v) for h, v in zip(headers, values) if h}, ensure_ascii=False)
            records.append(rec)
    finally:
        wb.close()
    if not records:
        raise ValueError('No se encontraron registros; no se genera una carga vacía.')
    locations = sorted({tuple(r[f] for f in ('departamento', 'provincia', 'distrito')) for r in records})
    ids = {loc: n for n, loc in enumerate(locations, 1)}
    for r in records:
        r['ubicacion_id'] = ids[tuple(r[f] for f in ('departamento', 'provincia', 'distrito'))]
    (output / 'data').mkdir(parents=True, exist_ok=True)
    (output / 'sql').mkdir(parents=True, exist_ok=True)
    with (output / 'data/bibliotecas.csv').open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=['registro_id', *FIELDS, 'ubicacion_id', 'hoja_origen', 'fila_origen'])
        writer.writeheader()
        writer.writerows({key: r[key] for key in writer.fieldnames} for r in records)
    sql = ['-- Generado a partir del Excel oficial; reemplaza la instantánea completa.',
           'SET XACT_ABORT ON;', 'BEGIN TRANSACTION;', 'DELETE FROM dbo.biblioteca;', 'DELETE FROM dbo.ubicacion;']
    sql += insert_lines('ubicacion', ['ubicacion_id', 'departamento', 'provincia', 'distrito'],
                        [(ids[loc], *loc) for loc in locations])
    bcols = ['registro_id', 'nombre_biblioteca', 'tipo_biblioteca', 'entidad_responsable', 'ubicacion_id', 'hoja_origen', 'fila_origen', 'datos_originales']
    sql += insert_lines('biblioteca', bcols, [tuple(r[c] for c in bcols) for r in records])
    sql.append('COMMIT;')
    (output / 'sql/02_load_data.sql').write_text('\n\n'.join(sql) + '\n', encoding='utf-8')
    metadata = dict(source_file=excel.name, sha256=source_sha, sheet=cfg['sheet'], header_row=header_row,
                    registros=len(records), ubicaciones=len(locations), columns=cols)
    (output / 'data/provenance.json').write_text(json.dumps(metadata, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f"Preparados {len(records)} registros y {len(locations)} ubicaciones. No se ha deduplicado por nombre ni imputado valores numéricos.")

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--excel', type=Path, default=ROOT / 'data/raw/bibliotecas.xlsx')
    parser.add_argument('--mapping', type=Path, default=ROOT / 'data/mapping.json')
    parser.add_argument('--output', type=Path, default=ROOT)
    parser.add_argument('--inspect', action='store_true')
    args = parser.parse_args()
    if not args.excel.exists():
        parser.error(f'Falta {args.excel}; descarga el Excel oficial.')
    if args.inspect:
        wb = openpyxl.load_workbook(args.excel, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                print(f'\nHOJA: {ws.title}')
                for n, row in enumerate(ws.iter_rows(max_row=10, values_only=True), 1):
                    print(n, [clean(v) for v in row])
        finally:
            wb.close()
    else:
        prepare(args.excel, args.mapping, args.output)

if __name__ == '__main__':
    main()
